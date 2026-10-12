"""Build a styled Excel workbook with claim-level data and live Excel formulas."""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Claim


REPORT_HEADERS = (
    "Claim ID", "Policy Number", "Claimant Name", "Claim Type",
    "Incident Date", "Reported Date", "Status", "Amount GBP",
    "Postcode", "Description",
)


def excel_safe(value: str) -> str:
    """Prevent untrusted text from being interpreted as an Excel formula."""
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def style_header(sheet, row: int, columns: int) -> None:
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 24


def create_claims_workbook(session: Session) -> BytesIO:
    claims = session.scalars(select(Claim).order_by(Claim.claim_id)).all()
    wb = Workbook()
    details = wb.active
    details.title = "Claims"
    details.append(REPORT_HEADERS)
    for claim in claims:
        details.append((
            excel_safe(claim.claim_id), excel_safe(claim.policy_number),
            excel_safe(claim.claimant_name), claim.claim_type,
            claim.incident_date, claim.reported_date, claim.status,
            float(claim.amount_gbp), excel_safe(claim.postcode), excel_safe(claim.description),
        ))
    style_header(details, 1, len(REPORT_HEADERS))
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:J{max(len(claims) + 1, 2)}"
    for index, width in enumerate((18, 20, 25, 16, 18, 18, 18, 19, 13, 42), 1):
        details.column_dimensions[get_column_letter(index)].width = width
    for row in details.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "dd mmm yyyy"
    for cell in details["H"][1:]:
        cell.number_format = '"£"#,##0.00'

    last_row = max(2, len(claims) + 1)
    lookup_range = f"Claims!$A$2:$A${last_row}"
    name_range = f"Claims!$C$2:$C${last_row}"
    status_range = f"Claims!$G$2:$G${last_row}"
    amount_range = f"Claims!$H$2:$H${last_row}"

    summary = wb.create_sheet("Summary")
    summary.append(("Northstar Claims Ltd — Claims Report", "Value"))
    summary.append(("Total claims", f"=COUNTA({lookup_range})"))
    summary.append(("Total claimed (GBP)", f"=SUM({amount_range})"))
    summary.append(("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")))
    summary["A8"] = "Status"
    summary["B8"] = "Claim count"
    summary["C8"] = "Total (GBP)"
    for row, status in enumerate(("Open", "In Progress", "Closed", "Declined"), start=9):
        summary.cell(row, 1, status)
        summary.cell(row, 2, f"=COUNTIF({status_range},A{row})")
        summary.cell(row, 3, f"=SUMIFS({amount_range},{status_range},A{row})")
        summary.cell(row, 3).number_format = '"£"#,##0.00'
    summary["B3"].number_format = '"£"#,##0.00'
    summary.column_dimensions["A"].width = 45
    summary.column_dimensions["B"].width = 24
    summary.column_dimensions["C"].width = 22
    style_header(summary, 1, 2)
    style_header(summary, 8, 3)

    lookup = wb.create_sheet("Claim Lookup")
    lookup["A1"] = "Claim Lookup (Microsoft 365 / Excel 2021+)"
    lookup["A2"], lookup["B2"] = "Claim ID (edit this)", claims[0].claim_id if claims else ""
    lookup["A3"] = "Claimant"
    lookup["A4"] = "Status"
    lookup["A5"] = "Amount GBP"
    lookup["B3"] = f'=_xlfn.XLOOKUP(B2,{lookup_range},{name_range},"Not found")'
    lookup["B4"] = f'=_xlfn.XLOOKUP(B2,{lookup_range},{status_range},"Not found")'
    lookup["B5"] = f'=_xlfn.XLOOKUP(B2,{lookup_range},{amount_range},"Not found")'
    lookup["B5"].number_format = '"£"#,##0.00'
    lookup.column_dimensions["A"].width = 43
    lookup.column_dimensions["B"].width = 36
    style_header(lookup, 1, 2)

    # openpyxl writes formulas but cannot calculate them; Excel recalculates on open.
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
