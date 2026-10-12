"""Three Task 2 Excel exports, following Task 1 workbook/header styling."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Asset, Ticket, User
from app.services.helpdesk_service import sla_status
from app.services.report_service import excel_safe, style_header


def finish(wb: Workbook) -> BytesIO:
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        style_header(sheet, 1, sheet.max_column)
        for idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(idx)].width = 23 if idx < 7 else 29
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def tickets_report(db: Session) -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Tickets'
    sheet.append(['Reference','Title','Employee','Technician','Category','Priority','Status',
                  'SLA','Created UTC','SLA due UTC','Resolved UTC','Resolution','Asset tag'])
    for t in db.scalars(select(Ticket).order_by(Ticket.id)):
        sheet.append([t.reference, excel_safe(t.title), excel_safe(t.employee.full_name),
                      excel_safe(t.technician.full_name) if t.technician else '', t.category,
                      t.priority, t.status, sla_status(t), t.created_at, t.sla_due_at,
                      t.resolved_at, excel_safe(t.resolution_details or ''),
                      excel_safe(t.asset.asset_tag) if t.asset else ''])
    summary = wb.create_sheet('Summary')
    summary.append(['Metric', 'Value'])
    summary.append(['Open / unresolved', '=COUNTIF(Tickets!G:G,"Open")+COUNTIF(Tickets!G:G,"Assigned")+COUNTIF(Tickets!G:G,"In Progress")'])
    summary.append(['Resolved', '=COUNTIF(Tickets!G:G,"Resolved")'])
    summary.append(['Critical', '=COUNTIF(Tickets!F:F,"Critical")'])
    return finish(wb)


def inventory_report(db: Session) -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Inventory'
    sheet.append(['Asset Tag','Serial','Device','Manufacturer','Model','OS','Status',
                  'Employee','Purchase date','Warranty expiry','Notes'])
    for a in db.scalars(select(Asset).order_by(Asset.asset_tag)):
        sheet.append([excel_safe(a.asset_tag), excel_safe(a.serial_number), a.device_type,
                      excel_safe(a.manufacturer), excel_safe(a.model),
                      excel_safe(a.operating_system or ''), a.status,
                      excel_safe(a.assigned_employee.full_name) if a.assigned_employee else '',
                      a.purchase_date, a.warranty_expiry, excel_safe(a.notes)])
    summary = wb.create_sheet('Summary')
    summary.append(['Metric', 'Value'])
    summary.append(['Assigned devices', '=COUNTIF(Inventory!G:G,"Assigned")'])
    summary.append(['Maintenance devices', '=COUNTIF(Inventory!G:G,"Maintenance")'])
    return finish(wb)


def workload_report(db: Session) -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Workload'
    sheet.append(['Technician', 'Assigned tickets', 'In Progress', 'Resolved', 'Closed', 'Open workload'])
    technicians = db.scalars(select(User).where(User.role == 'Technician').order_by(User.id)).all()
    for t in technicians:
        assigned = db.scalars(select(Ticket).where(Ticket.assigned_technician_id == t.id)).all()
        sheet.append([excel_safe(t.full_name), len(assigned),
                      sum(x.status == 'In Progress' for x in assigned),
                      sum(x.status == 'Resolved' for x in assigned),
                      sum(x.status == 'Closed' for x in assigned),
                      sum(x.status not in {'Resolved','Closed'} for x in assigned)])
    return finish(wb)
