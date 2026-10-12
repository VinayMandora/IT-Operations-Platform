"""Create exactly 1,000 entirely fictional UK claims as .xlsx and .csv."""
import csv
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.validation import HEADERS

OUTPUT_DIR = Path(__file__).resolve().parents[1] / "sample_data"
STATUSES = ("Open", "In Progress", "Closed", "Declined")
CLAIM_TYPES = ("Motor", "Home", "Travel", "Business")
POSTCODES = ("M1 1AA", "M2 3AB", "M3 4AD", "M4 5AF", "M5 6AG")


def generate_rows() -> list[dict]:
    rows = []
    for i in range(1, 901):
        incident = date(2024, 1, 1) + timedelta(days=i % 300)
        reported = incident + timedelta(days=(i % 9) + 1)
        rows.append({
            "claim_id": f"{' ncl' if i % 11 == 0 else 'NCL'}-{i:06d}{' ' if i % 17 == 0 else ''}",
            "policy_number": f"pol-{i:06d}" if i % 13 == 0 else f"POL-{i:06d}",
            "claimant_name": f"  test   claimant {i:04d} " if i % 7 == 0 else f"Test Claimant {i:04d}",
            "claim_type": CLAIM_TYPES[i % len(CLAIM_TYPES)].lower() if i % 19 == 0 else CLAIM_TYPES[i % len(CLAIM_TYPES)],
            "incident_date": incident,
            "reported_date": reported,
            "status": f" {STATUSES[i % 4].lower()} " if i % 9 == 0 else STATUSES[i % 4],
            "amount_gbp": round(100 + i * 0.75, 2),
            "postcode": POSTCODES[i % len(POSTCODES)].lower() if i % 8 == 0 else POSTCODES[i % len(POSTCODES)],
            "description": f"Fictional portfolio test claim {i:04d}",
        })
    for i in range(1, 41):
        duplicate = rows[i - 1].copy()
        duplicate["description"] = "INTENTIONAL DUPLICATE: " + duplicate["description"]
        rows.append(duplicate)
    # 60 intentionally bad records: five groups of 12.
    for i in range(60):
        item = rows[(i + 100) % 900].copy()
        item["claim_id"] = f"NCL-{901 + i:06d}"
        if i < 12:
            item["claim_id"] = "INVALID-ID"
        elif i < 24:
            item["reported_date"] = "31/02/2025"
        elif i < 36:
            item["amount_gbp"] = "-50.00"
        elif i < 48:
            item["status"] = "Pending unknown"
        else:
            item["postcode"] = "NOT A POSTCODE"
        item["description"] = "INTENTIONAL INVALID: " + item["description"]
        rows.append(item)
    assert len(rows) == 1000
    return rows


def main() -> None:
    OUTPUT_DIR.mkdir(exist_ok=True, parents=True)
    rows = generate_rows()
    xlsx_path = OUTPUT_DIR / "northstar_claims_1000.xlsx"
    csv_path = OUTPUT_DIR / "northstar_claims_1000.csv"

    wb = Workbook()
    ws = wb.active
    ws.title = "ClaimsData"
    ws.append(HEADERS)
    for i, row in enumerate(rows, start=2):
        ws.append([row[header] for header in HEADERS])
        if i >= 902:
            for cell in ws[i]:
                cell.fill = PatternFill("solid", fgColor="FFF0E8")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J1001"
    for column, width in {"A": 18, "B": 18, "C": 25, "D": 16, "E": 18, "F": 18, "G": 16, "H": 16, "I": 16, "J": 50}.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    for cell in ws["H"][1:]:
        if isinstance(cell.value, (float, int)):
            cell.number_format = '"£"#,##0.00'
    wb.save(xlsx_path)

    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    print("Created 1,000 rows: 900 valid unique, 40 duplicate, 60 invalid")
    print(xlsx_path)
    print(csv_path)


if __name__ == "__main__":
    main()
