"""Inventory permissions, duplicate checks, conflicts, return history and Excel exports."""
from datetime import date
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.models import AssetMovement, AuditEvent
from tests.conftest import login_as


def laptop():
    return {'asset_tag':'NS-LAP-NEW','serial_number':'SERIAL-100','device_type':'Laptop',
            'manufacturer':'Dell', 'model':'Latitude 5440','operating_system':'Windows 11',
            'purchase_date':'2025-01-02','warranty_expiry':'2028-01-01','notes':'Spare'}


def test_asset_permissions_and_validation(client):
    emp = login_as(client,'employee.ellie')
    assert client.get('/api/assets').status_code == 403
    assert client.post('/api/assets',json=laptop(),headers=emp).status_code == 403
    admin = login_as(client)
    assert client.post('/api/assets',json={**laptop(),'warranty_expiry':'2020-01-01'},headers=admin).status_code == 422
    response=client.post('/api/assets',json=laptop(),headers=admin)
    assert response.status_code == 201,response.text
    assert client.post('/api/assets',json=laptop(),headers=admin).status_code == 409
    assert len(client.get('/api/assets?q=Latitude').json()['items']) == 1


def test_device_assigned_and_returned_without_conflict(client,session_factory):
    tech=login_as(client,'tech.owen')
    response=client.post('/api/assets',json=laptop(),headers=tech)
    assert response.status_code == 201,response.text
    aid=response.json()['id']
    first=client.post(f'/api/assets/{aid}/assign',json={'employee_id':4},headers=tech)
    assert first.status_code == 200, first.text
    assert first.json()['assigned_employee_id']==4
    assert client.post(f'/api/assets/{aid}/assign',json={'employee_id':5},headers=tech).status_code==409
    assert client.patch(f'/api/assets/{aid}',json={'status':'Maintenance'},headers=tech).status_code==409
    assert client.get('/api/assets?status=Assigned').json()['items'][0]['id']==aid
    assert client.post(f'/api/assets/{aid}/return',headers=tech).json()['status']=='Available'
    assert client.post(f'/api/assets/{aid}/return',headers=tech).status_code==409
    patched=client.patch(f'/api/assets/{aid}',json={'status':'Maintenance','notes':'Fan cleaning due'},headers=tech)
    assert patched.status_code == 200 and patched.json()['status']=='Maintenance'
    assert client.post(f'/api/assets/{aid}/assign',json={'employee_id':5},headers=tech).status_code==409
    with session_factory() as db:
        assert db.scalar(select(func.count()).select_from(AssetMovement)) == 2
        assert db.scalar(select(func.count()).select_from(AuditEvent).where(AuditEvent.event_type=='ASSET_RETURNED')) == 1


def test_ticket_asset_employee_ownership_and_exports(client):
    tech=login_as(client,'tech.owen')
    aid=client.post('/api/assets',json=laptop(),headers=tech).json()['id']
    assert client.post(f'/api/assets/{aid}/assign',json={'employee_id':4},headers=tech).status_code == 200
    employee=login_as(client,'employee.jamie')
    from tests.test_helpdesk import payload
    assert client.post('/api/tickets',json=payload(aid),headers=employee).status_code==403
    employee=login_as(client,'employee.ellie')
    created=client.post('/api/tickets',json=payload(aid),headers=employee)
    assert created.status_code==201 and created.json()['asset_tag']=='NS-LAP-NEW'
    assert client.get('/api/reports/inventory.xlsx').status_code==403
    tech=login_as(client,'tech.owen')
    for kind, name in [('tickets','Tickets'),('inventory','Inventory')]:
        response=client.get(f'/api/reports/{kind}.xlsx')
        assert response.status_code==200
        workbook=load_workbook(BytesIO(response.content))
        assert name in workbook.sheetnames and 'Summary' in workbook.sheetnames
    assert client.get('/api/reports/workload.xlsx').status_code==403
    admin=login_as(client)
    assert client.get('/api/reports/workload.xlsx').status_code==200


def test_html_pages_role_visibility(client):
    admin=login_as(client,'admin')
    assert client.get('/assets').status_code==200
    assert client.get('/users').status_code==200
    assert 'Imported claims' in client.get('/').text
    login_as(client,'employee.ellie')
    assert client.get('/users').status_code==403
    assert client.get('/assets').status_code==403
    assert client.get('/tickets/new').status_code==200
    assert 'Imported claims' not in client.get('/').text
    assert client.get('/api/claims').status_code==403
