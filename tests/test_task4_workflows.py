"""Task 4 end-to-end tests; isolated SQLite, no Zoho credentials."""
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import (Claim, ClaimCRMState, CommunicationTemplate, CRMOperation,
                        Ticket, WorkflowExecution, WorkflowRule, WorkflowTransition)
from app.services.claims_workflow import seed_rules
from app.services.claim_templates import seed_templates
from app.services.zoho_client import MockZohoClient, CRMError
from app.routers.claims_crm import get_zoho
from scripts.generate_sample_data import generate_rows
from tests.conftest import login_as
from tests.test_import_service import one_csv


def setup_claim(client, maker, reference='NCL-000001'):
    headers = login_as(client)
    row = next(r for r in generate_rows() if r['claim_id'] == reference)
    response = client.post('/api/claims/import', headers=headers,
                           files={'file': ('sample.csv', one_csv(row), 'text/csv')})
    assert response.status_code == 200, response.text
    with maker() as db:
        seed_rules(db)
        seed_templates(db)
    return headers


def send_event(client, headers, event, key='event-12345678', claim='NCL-000001'):
    return client.post('/api/workflows/events', headers=headers,
        json={'claim_id':claim,'event_type':event,'event_key':key})


def test_full_lifecycle_deduplication_and_report(client, session_factory):
    hdr = setup_claim(client, session_factory)
    stages = [('submitted','Assigned'), ('missing_documents','Evidence Required'),
              ('evidence_received','Ready for Review'), ('resolved','Resolved')]
    for i,(event,expected) in enumerate(stages):
        result = send_event(client, hdr, event, f'lifecycle-{i:04}').json()
        assert result['results'][0]['status'] == 'Success', result
        assert client.get('/api/crm/claims/NCL-000001').json()['stage'] == expected
        repeated = send_event(client, hdr, event, f'lifecycle-{i:04}').json()
        assert repeated['results'][0]['replayed'] is True
    last = client.get('/api/crm/claims/NCL-000001').json()
    assert last['claim_status'] == 'Closed'
    assert last['missing_documentation'] is False
    assert last['completed_at']
    with session_factory() as db:
        assert db.query(WorkflowExecution).count() == 4
        assert db.query(WorkflowTransition).count() == 4
        assert db.query(Claim).first().status == 'Closed'
    output = client.get('/api/reports/claims-workflows.xlsx')
    assert output.status_code == 200
    wb = load_workbook(BytesIO(output.content))
    assert wb.sheetnames == ['Summary','Claims by Status','Status Summary','Missing Documents',
                             'Failed Workflows','Processing Time','Intervention']
    assert 'AVERAGE' in wb['Summary']['B5'].value
    assert wb['Claims by Status']['A2'].value == 'NCL-000001'


def test_failed_workflow_helpdesk_and_admin_retry(client, session_factory):
    hdr = setup_claim(client, session_factory)
    with session_factory() as db:
        db.add(WorkflowRule(name='Broken example', event_type='submitted',
                            action_type='demo_fail', parameters_json='{}', enabled=True))
        db.commit()
    result = send_event(client, hdr, 'submitted', 'broken-event-111').json()
    assert [v['status'] for v in result['results']] == ['Success','Failed']
    failed = result['results'][1]
    assert failed['support_ticket_id']
    with session_factory() as db:
        assert db.query(Ticket).count() == 1
    duplicate = send_event(client, hdr, 'submitted', 'broken-event-111').json()
    assert duplicate['results'][1]['replayed'] is True
    assert client.get('/api/workflows/executions').json()['items'][0]['status'] == 'Failed'
    with session_factory() as db:
        rule = db.scalar(select(WorkflowRule).where(WorkflowRule.name == 'Broken example'))
        rule.action_type = 'ready_for_review'
        db.commit()
    retry = client.post(f"/api/workflows/executions/{failed['execution_id']}/retry", headers=hdr)
    assert retry.status_code == 200
    assert retry.json()['status'] == 'Failed'  # Cannot review evidence that has not been requested
    assert retry.json()['retry_count'] == 1
    with session_factory() as db:
        rule = db.scalar(select(WorkflowRule).where(WorkflowRule.name == 'Broken example'))
        rule.action_type = 'evidence_required'
        db.commit()
    resolved = client.post(f"/api/workflows/executions/{failed['execution_id']}/retry",headers=hdr)
    assert resolved.json()['status'] == 'Success'
    assert resolved.json()['retry_count'] == 2
    with session_factory() as db:
        assert db.query(Ticket).count() == 1  # Never duplicate the escalation ticket


def test_mock_crm_pull_push_and_retry(client, session_factory):
    hdr = setup_claim(client, session_factory)
    mocked = MockZohoClient(records=[{'id':'mock-1001','Claim_Reference':'NCL-000001',
                                     'Status':'Open', 'Missing_Documents':False}])
    client.app.dependency_overrides[get_zoho] = lambda: mocked
    pulled = client.post('/api/crm/claims/NCL-000001/sync', headers=hdr,
                         json={'operation_key':'pull-test-2026'})
    assert pulled.status_code == 200, pulled.text
    assert pulled.json()['status'] == 'Success'
    assert client.post('/api/crm/claims/NCL-000001/sync',headers=hdr,
        json={'operation_key':'pull-test-2026'}).json()['replayed'] is True
    assert mocked.get_claim('mock-1001')['Status'] == 'Open'
    assert send_event(client,hdr,'submitted','crm-new-event1').json()['results'][0]['status'] == 'Success'
    pushed = client.post('/api/crm/claims/NCL-000001/push', headers=hdr,
                          json={'operation_key':'push-test-2026'})
    assert pushed.json()['status'] == 'Success'
    assert mocked.get_claim('mock-1001')['Status'] == 'Assigned'

    # A failed pull records its reason and can be retried after the CRM record appears.
    missing = MockZohoClient(records=[])
    client.app.dependency_overrides[get_zoho] = lambda: missing
    failed = client.post('/api/crm/claims/NCL-000001/sync', headers=hdr,
                        json={'operation_key':'failed-pull-111'}).json()
    assert failed['status'] == 'Failed'
    assert failed['retryable'] is True
    assert client.post('/api/crm/claims/NCL-000001/sync', headers=hdr,
                      json={'operation_key':'failed-pull-111'}).json()['replayed'] is True
    missing.records = dict(mocked.records)
    retry = client.post(f"/api/crm/operations/{failed['operation_id']}/retry", headers=hdr)
    assert retry.json()['status'] == 'Success'
    assert retry.json()['retry_count'] == 1


def test_role_checks_and_csrf(client, session_factory):
    hdr = setup_claim(client, session_factory)
    assert client.post('/api/workflows/events', json={'claim_id':'NCL-000001',
        'event_type':'submitted','event_key':'auth-event-111'}).status_code == 403
    login_as(client,'employee.ellie')
    assert client.get('/claims-management').status_code == 403
    assert client.get('/workflow-logs').status_code == 403
    assert client.get('/communication-templates').status_code == 403
    assert client.get('/api/crm/claims').status_code == 403
    assert client.get('/api/reports/claims-workflows.xlsx').status_code == 403
    assert send_event(client, {}, 'submitted').status_code == 403


def test_template_edit_preview_and_injection_rejected(client, session_factory):
    hdr = setup_claim(client, session_factory)
    route = '/api/workflows/templates/acknowledgement'
    assert client.put(route, headers=hdr, json={'subject':'{claim_id} update',
          'body':'Dear {claimant_name}, {claim_status}'}).status_code == 200
    preview = client.get(route + '/preview/NCL-000001').json()
    assert preview['subject'] == 'NCL-000001 update' and preview['sent'] is False
    rejected = client.put(route, headers=hdr, json={'subject':'{claimant_name.__class__}',
          'body':'Safe'})
    assert rejected.status_code == 422
    assert client.put(route,headers=hdr,json={'subject':'hello {claim_id}',
          'body':'hello {claimant_name[0]}'}).status_code == 422
    assert client.get('/communication-templates').status_code == 200


def test_rule_configuration_and_pages(client, session_factory):
    hdr = setup_claim(client, session_factory)
    assert client.get('/claims-management').status_code == 200
    assert client.get('/workflow-logs').status_code == 200
    duplicate = client.post('/api/workflows/rules', headers=hdr,
        json={'name':'Assign new claim to claims team','event_type':'submitted','action_type':'assign_team'})
    assert duplicate.status_code == 409
    added = client.post('/api/workflows/rules', headers=hdr,
        json={'name':'Training follow-up','event_type':'submitted','action_type':'assign_team',
              'handler':'Claims Team'})
    assert added.status_code == 201
    rule_id = added.json()['id']
    disabled = client.patch(f'/api/workflows/rules/{rule_id}', headers=hdr, json={'enabled':False})
    assert disabled.json()['enabled'] is False
    assert client.get('/api/workflows/rules').status_code == 200
    assert client.get('/api/workflows/transitions').status_code == 200


def test_two_distinct_unsynced_states_and_intervention(client, session_factory):
    """SQL Server filtered unique index must allow multiple unlinked Zoho IDs."""
    from pathlib import Path
    hdr = login_as(client)
    data = (Path(__file__).resolve().parents[1] / 'sample_data' / 'task4_demo_claims.csv').read_bytes()
    response = client.post('/api/claims/import', files={'file':('task4_demo_claims.csv',data,'text/csv')}, headers=hdr)
    assert response.json()['imported_rows'] == 2
    with session_factory() as db:
        seed_rules(db)
    for reference in ['NCL-009999','NCL-009998']:
        assert send_event(client,hdr,'submitted','new-'+reference,reference).json()['results'][0]['status'] == 'Success'
    assert client.get('/api/crm/claims').json()['count'] == 2
    with session_factory() as db:
        assert db.query(ClaimCRMState).count() == 2
        assert all(s.zoho_record_id is None for s in db.query(ClaimCRMState))
    bad = send_event(client,hdr,'evidence_received','bad-evidence-999','NCL-009998').json()['results'][0]
    assert bad['status'] == 'Failed'
    excel = load_workbook(BytesIO(client.get('/api/reports/claims-workflows.xlsx').content))
    assert excel['Intervention']['A2'].value == 'NCL-009998'
