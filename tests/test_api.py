"""End-to-end FastAPI tests without a running SQL Server instance."""
from io import BytesIO

from openpyxl import load_workbook

from scripts.generate_sample_data import generate_rows
from tests.test_import_service import one_csv
from tests.conftest import login_as


def test_health_and_dashboard(client):
    assert client.get("/health").json()["status"] == "ok"
    response = client.get("/")
    assert response.status_code == 200
    assert "Northstar Claims Ltd" in response.text


def test_import_list_and_formula_report(client):
    headers = login_as(client)
    csv_bytes = one_csv(generate_rows()[0])
    upload = client.post("/api/claims/import", files={"file": ("claims.csv", csv_bytes, "text/csv")}, headers=headers)
    assert upload.status_code == 200, upload.text
    batch = upload.json()
    assert batch["imported_rows"] == 1
    claims = client.get("/api/claims?limit=10").json()
    assert claims["total"] == 1
    assert claims["items"][0]["claim_id"] == "NCL-000001"
    imports = client.get("/api/imports").json()
    assert imports["items"][0]["batch_id"] == batch["batch_id"]
    report = client.get("/api/reports/claims.xlsx")
    assert report.status_code == 200
    workbook = load_workbook(BytesIO(report.content), data_only=False)
    assert workbook.sheetnames == ["Claims", "Summary", "Claim Lookup"]
    assert workbook["Claims"]["A2"].value == "NCL-000001"
    assert "COUNTIF" in workbook["Summary"]["B9"].value
    assert "SUMIFS" in workbook["Summary"]["C9"].value
    assert "XLOOKUP" in workbook["Claim Lookup"]["B3"].value


def test_rejections_visible_in_api(client):
    headers = login_as(client)
    row = generate_rows()[0].copy()
    row["status"] = "INVALID"
    response = client.post("/api/claims/import", files={"file": ("bad.csv", one_csv(row), "text/csv")}, headers=headers)
    assert response.status_code == 200
    assert response.json()["rejected_rows"] == 1
    batch_id = response.json()["batch_id"]
    details = client.get(f"/api/imports/{batch_id}/rejections").json()
    assert "status must" in details["items"][0]["reason"]


def test_invalid_extension_and_oversized_upload(client, monkeypatch):
    from app.routers import claims
    headers = login_as(client)
    wrong_type = client.post("/api/claims/import", files={"file": ("test.txt", b"hello", "text/plain")}, headers=headers)
    assert wrong_type.status_code == 400
    monkeypatch.setattr(claims, "MAX_UPLOAD_BYTES", 10)
    oversized = client.post("/api/claims/import", files={"file": ("test.csv", b"x" * 11, "text/csv")}, headers=headers)
    assert oversized.status_code == 413


def test_formula_like_user_text_exported_as_text():
    from app.services.report_service import excel_safe
    assert excel_safe("=2+2") == "'=2+2"
    assert excel_safe("@SUM(A1:A3)") == "'@SUM(A1:A3)"
    assert excel_safe("ordinary") == "ordinary"
